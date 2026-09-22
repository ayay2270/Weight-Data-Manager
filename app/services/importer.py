from __future__ import annotations
import hashlib, json, re
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from ..models import ImportBatch, Measurement
from ..schemas import MeasurementInput
from .measurements import create_measurement


SHEETS = {"Part Level": ("Part", "g", 1), "Node Level": ("Node", "kg", 1), "Rack Level": ("Rack", "kg", 1), "Package": ("Package", "kg", 2)}


def clean(value):
    if value is None: return None
    if isinstance(value, str):
        value = value.replace("\u200b", "").replace("\xa0", " ").strip()
    return value


def text(value):
    value = clean(value)
    if value is None or str(value).upper() in {"", "NA", "N/A"}: return None
    return str(value).strip()


def parse_weight(raw):
    value = clean(raw)
    if value is None: return None, "TBD", None
    if isinstance(value, (int, float)) and float(value) > 0: return float(value), "Measured", None
    try:
        number = float(str(value))
        if number > 0: return number, "Measured", None
    except ValueError:
        pass
    note = f"Original weight value requires review: {value}"
    return None, "TBD", note


def merged_master_value(ws, coordinate):
    for rng in ws.merged_cells.ranges:
        if coordinate in rng:
            return ws.cell(rng.min_row, rng.min_col).value, str(rng)
    return ws[coordinate].value, None


def row_fingerprint(sheet, row, values):
    payload = json.dumps([sheet, row] + [str(clean(v)) for v in values], ensure_ascii=False)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def import_excel(db: Session, path: str | Path, user="Excel Import"):
    path = Path(path)
    file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    existing = db.scalar(select(ImportBatch).where(ImportBatch.file_hash == file_hash))
    if existing:
        result = json.loads(existing.result_json)
        result["duplicate_file"] = True
        return result
    wb = load_workbook(path, data_only=True)
    result = {"successfully_imported": 0, "needs_review": 0, "skipped": 0, "errors": 0, "details": [], "duplicate_file": False}
    for sheet_name, (level, unit, header_row) in SHEETS.items():
        if sheet_name not in wb.sheetnames:
            result["errors"] += 1; result["details"].append(f"Missing sheet: {sheet_name}"); continue
        ws = wb[sheet_name]
        headers = [clean(ws.cell(header_row, c).value) for c in range(1, ws.max_column + 1)]
        index = {str(v): i for i, v in enumerate(headers) if v}
        for row in range(header_row + 1, ws.max_row + 1):
            values = [ws.cell(row, c).value for c in range(1, ws.max_column + 1)]
            if not any(clean(v) not in (None, "") for v in values): continue
            try:
                def get(name): return values[index[name]] if name in index else None
                raw_weight, merged = merged_master_value(ws, f"{chr(65 + index[next(k for k in index if k.startswith('Weight'))])}{row}")
                weight, value_type, weight_note = parse_weight(raw_weight)
                project = text(get("Project")) or "Needs Review"
                description = text(get("Part Description"))
                if not description:
                    result["skipped"] += 1; result["details"].append(f"{sheet_name} row {row}: missing description"); continue
                import_notes = [n for n in [weight_note, f"Merged source cell {merged}; interpreted as combined weight" if merged else None, "Missing project in source" if project == "Needs Review" else None] if n]
                note = text(get("Note"))
                if import_notes: note = " | ".join([x for x in [note, *import_notes] if x])
                date_value = get("Measured Date")
                measured_date = date_value.date() if isinstance(date_value, datetime) else date_value
                data = MeasurementInput(project=project, level=level, description=description, lenovo_pn=text(get("Lenovo PN")), customer_pn=text(get("MSFT PN")), manufacturer=text(get("Manufacturer")), category=text(get("Part Category")) or ("Package" if level == "Package" else None), weight_value=weight, weight_unit=unit, value_type=value_type, measured_date=measured_date, measured_by=user, project_phase=None, data_source="Internal", note=note, package_type="E2010 48U Rack" if level == "Package" else None, applicable_parent="E2010 48U Rack" if level == "Package" else None, combined_weight=bool(merged))
                fingerprint = row_fingerprint(sheet_name, row, values)
                if db.scalar(select(Measurement.id).where(Measurement.import_fingerprint == fingerprint)):
                    result["skipped"] += 1; continue
                needs_review = value_type == "TBD" or project == "Needs Review" or bool(merged)
                create_measurement(db, data, user, status="Draft" if needs_review else "Approved", import_fingerprint=fingerprint, original_value=str(clean(raw_weight)) if raw_weight is not None else None, import_note="; ".join(import_notes) or None)
                result["successfully_imported"] += 1
                if needs_review: result["needs_review"] += 1
            except Exception as exc:
                db.rollback(); result["errors"] += 1; result["details"].append(f"{sheet_name} row {row}: {exc}")
    db.add(ImportBatch(file_hash=file_hash, filename=path.name, imported_by=user, result_json=json.dumps(result, ensure_ascii=False)))
    db.commit()
    return result

